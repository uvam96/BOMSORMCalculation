"""Core logic: parse SAP ByD analytics CSV exports and explode SO demand
through multi-level BOMs (FG -> SFG -> ... -> RM)."""
import io
import re
import pandas as pd


# ---------------------------------------------------------------- parsing
def _read_csv(file) -> pd.DataFrame:
    """Read a SAP analytics CSV (UTF-8 BOM, CR-only line endings, blank headers)."""
    raw = file.read() if hasattr(file, "read") else open(file, "rb").read()
    if isinstance(raw, bytes):
        raw = raw.decode("utf-8-sig", errors="replace")
    raw = raw.replace("\r\n", "\n").replace("\r", "\n")
    df = pd.read_csv(io.StringIO(raw), dtype=str, keep_default_na=False)
    df.columns = [c.strip() for c in df.columns]
    return df


def _col(df, name):
    for c in df.columns:
        if c.lower() == name.lower():
            return c
    raise KeyError(f"Column '{name}' not found. Columns: {list(df.columns)}")


def _next_col(df, name):
    """SAP puts the description/ID in the blank-headed column right after a field."""
    i = list(df.columns).index(_col(df, name))
    return df.columns[i + 1]


_NUM = re.compile(r"^\s*(-?[\d,]*\.?\d+)\s*(.*)$")


def split_qty(s):
    """'1,234.50 3PCK' -> (1234.5, '3PCK')"""
    m = _NUM.match(str(s))
    if not m:
        return 0.0, ""
    return float(m.group(1).replace(",", "")), m.group(2).strip()


def parse_bom(files) -> pd.DataFrame:
    frames = []
    for f in files:
        d = _read_csv(f)
        out = pd.DataFrame({
            "Parent": d[_col(d, "Product ID")].str.strip(),
            "ParentDesc": d[_next_col(d, "Product ID")].str.strip(),
            "Group": d[_col(d, "Line Item Group ID")].str.strip() if "Line Item Group ID" in d.columns else "",
            "LineItem": d[_col(d, "Line Item ID")].str.strip() if "Line Item ID" in d.columns else "",
            "Comp": d[_col(d, "Raw Material")].str.strip(),
            "CompDesc": d[_next_col(d, "Raw Material")].str.strip(),
            "InQty": d[_col(d, "Input Quantity")],
            "OutQty": d[_col(d, "Output Quantity")],
        })
        frames.append(out)
    bom = pd.concat(frames, ignore_index=True)
    bom = bom[(bom.Parent != "") & (bom.Comp != "")]
    bom = bom.drop_duplicates(subset=["Parent", "Group", "LineItem", "Comp", "InQty", "OutQty"])
    bom[["QtyIn", "UoM"]] = bom.InQty.apply(lambda s: pd.Series(split_qty(s)))
    bom[["QtyOut", "OutUoM"]] = bom.OutQty.apply(lambda s: pd.Series(split_qty(s)))
    bom["QtyOut"] = bom.QtyOut.replace(0, 1.0)
    bom["QtyPer"] = bom.QtyIn / bom.QtyOut
    return bom.reset_index(drop=True)


def parse_so(file) -> pd.DataFrame:
    d = _read_csv(file)
    prod = _col(d, "Product")
    nxt = _next_col(d, "Product")
    # Product ID is the numeric one of the pair (SAP puts desc under 'Product')
    if d[prod].str.strip().str.fullmatch(r"\d+").all():
        pid, pdesc = d[prod], d[nxt]
    else:
        pid, pdesc = d[nxt], d[prod]
    so = pd.DataFrame({
        "Sales Order": d[_col(d, "Sales Order")].str.strip(),
        "SO Item": d[_col(d, "Sales Order (Item)")].str.strip(),
        "Customer PO": d[_col(d, "Customer PO")].str.strip() if "Customer PO" in d.columns else "",
        "External Reference": d[_col(d, "External Reference")].str.strip() if "External Reference" in d.columns else "",
        "FG ID": pid.str.strip(),
        "FG Description": pdesc.str.strip(),
    })
    req = d[_col(d, "Requested Quantity")].apply(split_qty)
    so["Requested Qty"] = [q for q, _ in req]
    so["UoM"] = [u for _, u in req]
    if "Fulfilled Quantity" in d.columns:
        so["Fulfilled Qty"] = d["Fulfilled Quantity"].apply(lambda s: split_qty(s)[0])
    else:
        so["Fulfilled Qty"] = 0.0
    so["Open Qty"] = (so["Requested Qty"] - so["Fulfilled Qty"]).clip(lower=0)
    if "List Price" in d.columns:
        so["List Price Value"] = d["List Price"].apply(lambda s: split_qty(s)[0])
    so = so[so["FG ID"] != ""]
    return so.reset_index(drop=True)


# ---------------------------------------------------------------- explosion
def explode(so: pd.DataFrame, bom: pd.DataFrame, qty_col="Requested Qty", max_depth=10):
    """Return (detail rows, missing-BOM list). Any component that has its own
    BOM in the uploaded data is exploded further; the rest are raw materials."""
    children = {p: g for p, g in bom.groupby("Parent")}
    rows, missing = [], set()

    def walk(so_row, item, desc, qty_per, path, depth):
        if depth > max_depth or item in path:
            raise ValueError(f"BOM loop or too deep at {' > '.join(path + [item])}")
        for _, c in children[item].iterrows():
            q = qty_per * c.QtyPer
            if c.Comp in children:
                walk(so_row, c.Comp, c.CompDesc, q, path + [item], depth + 1)
            else:
                rows.append({
                    "Sales Order": so_row["Sales Order"], "SO Item": so_row["SO Item"],
                    "FG ID": so_row["FG ID"], "FG Description": so_row["FG Description"],
                    "Parent ID": item, "Parent Description": desc,
                    "Level": depth + 1,
                    "Path": " > ".join(path + [item]),
                    "Line Group": re.sub(r"^\d+\.", "", c.Group or ""),
                    "RM ID": c.Comp, "RM Description": c.CompDesc, "UoM": c.UoM,
                    "RM Qty per FG": q,
                    "SO Qty": so_row[qty_col],
                    "RM Requirement": q * so_row[qty_col],
                })

    for _, r in so.iterrows():
        if r["FG ID"] not in children:
            missing.add((r["FG ID"], r["FG Description"]))
            continue
        walk(r, r["FG ID"], r["FG Description"], 1.0, [], 0)
    return pd.DataFrame(rows), sorted(missing)


def summarise(detail: pd.DataFrame, extra_waste_pct=0.0) -> pd.DataFrame:
    if detail.empty:
        return detail
    s = (detail.groupby(["RM ID", "RM Description", "UoM"], as_index=False)
         .agg(**{"Line Group": ("Line Group", lambda x: " / ".join(sorted(set(x)))),
                 "BOM Requirement": ("RM Requirement", "sum"),
                 "Used in FGs": ("FG ID", "nunique")}))
    s["Extra Allowance %"] = extra_waste_pct
    s["Total Requirement"] = s["BOM Requirement"] * (1 + extra_waste_pct / 100)
    return s.sort_values(["Line Group", "RM Description"]).reset_index(drop=True)


def size_matrix(detail: pd.DataFrame) -> pd.DataFrame:
    """RM x FG pivot for quick checking by size/colour."""
    if detail.empty:
        return detail
    return detail.pivot_table(index=["RM ID", "RM Description", "UoM"], columns="FG Description",
                              values="RM Requirement", aggfunc="sum", fill_value=0).reset_index()


# ---------------------------------------------------------------- export
def to_excel(summary, detail, so, bom, matrix, missing, meta) -> bytes:
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="RM Requirement", index=False, startrow=3)
        matrix.to_excel(xw, sheet_name="RM by FG", index=False)
        detail.to_excel(xw, sheet_name="BOM Explosion", index=False)
        so.to_excel(xw, sheet_name="Sales Order Lines", index=False)
        bom.drop(columns=["InQty", "OutQty"]).to_excel(xw, sheet_name="BOM Master", index=False)
        if missing:
            pd.DataFrame(missing, columns=["FG ID", "FG Description"]).to_excel(xw, sheet_name="Missing BOMs", index=False)
        ws = xw.sheets["RM Requirement"]
        ws["A1"] = "Raw Material Requirement"
        ws["A1"].font = Font(name="Arial", bold=True, size=14)
        ws["A2"] = meta
        ws["A2"].font = Font(name="Arial", italic=True, size=9, color="595959")
        for name, sh in xw.sheets.items():
            hrow = 4 if name == "RM Requirement" else 1
            for cell in sh[hrow]:
                cell.font = Font(name="Arial", bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F3864")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            sh.freeze_panes = sh.cell(hrow + 1, 1)
            for i, col in enumerate(sh.iter_cols(min_row=hrow, max_row=min(sh.max_row, 200)), 1):
                w = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                sh.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 9), 45)
                for c in col[1:]:
                    if isinstance(c.value, float):
                        c.number_format = "#,##0.0000"
    return buf.getvalue()
